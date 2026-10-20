"""StudyLens AI backend. Run: uvicorn backend.main:app --reload"""

import io
import os
import threading
import uuid
from datetime import datetime, timedelta, timezone
from typing import Optional

from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.staticfiles import StaticFiles
from pathlib import Path
from fastapi.concurrency import run_in_threadpool
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel, Field
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.utils import simpleSplit
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape

from backend.rag.groq_answer_service import generate_exam_answer, ProviderQuotaExceeded
from backend.rag.vector_store import (
    chunk_document,
    create_vector_index,
    search_vector_index,
)
from backend.services.bulk_answer_service import (
    create_rag_answer_function,
    flatten_questions,
)
from backend.services.pdf_extractor import extract_pdf_content
from backend.services.question_validator import validate_question_paper
from backend.services.vision_question_parser import parse_question_paper_with_vision

app = FastAPI(title="StudyLens AI", description="AI Document Intelligence Platform", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip() for origin in os.getenv(
        "FRONTEND_ORIGINS", "http://127.0.0.1:5500,http://localhost:5500"
    ).split(",") if origin.strip()],
    allow_credentials=False,
    allow_methods=["GET", "POST", "DELETE"],
    allow_headers=["*"],
)

MAX_UPLOAD_SIZE = 20 * 1024 * 1024
SESSION_TTL = timedelta(hours=2)
sessions_db = {}  # RAM only: all sessions disappear when the server restarts.
sessions_lock = threading.RLock()


def utcnow():
    return datetime.now(timezone.utc)


def purge_expired_sessions():
    cutoff = utcnow() - SESSION_TTL
    with sessions_lock:
        for key, session in list(sessions_db.items()):
            if session["last_activity"] < cutoff and not session.get("running"):
                del sessions_db[key]


def get_session(session_id: str):
    purge_expired_sessions()
    with sessions_lock:
        session = sessions_db.get(session_id)
        if session is None:
            raise HTTPException(status_code=404, detail="Session missing or expired. Upload textbook again.")
        session["last_activity"] = utcnow()
        return session


async def read_pdf(file: UploadFile):
    if not file.filename or not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Upload a .pdf file.")
    data = await file.read(MAX_UPLOAD_SIZE + 1)
    if len(data) > MAX_UPLOAD_SIZE:
        raise HTTPException(status_code=413, detail="PDF exceeds 20 MB.")
    if not data.startswith(b"%PDF-"):
        raise HTTPException(status_code=400, detail="File does not appear to be a PDF.")
    return data


class SearchRequest(BaseModel):
    question: str = Field(min_length=1)


class AskRequest(BaseModel):
    question: str = Field(min_length=1)
    allow_general: bool = False


@app.get("/api/health")
def health_check():
    return {"status": "ok"}


@app.post("/api/sessions/upload")
async def upload_pdf(file: UploadFile = File(...)):
    data = await read_pdf(file)
    try:
        extracted = await run_in_threadpool(extract_pdf_content, data)
        pages = extracted["pages"]
        chunks = await run_in_threadpool(chunk_document, pages)
        if not chunks:
            raise HTTPException(status_code=422, detail="No usable textbook text found.")
        index = await run_in_threadpool(create_vector_index, chunks)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not process textbook: {exc}") from exc

    session_id = str(uuid.uuid4())
    now = utcnow()
    with sessions_lock:
        sessions_db[session_id] = {
            "created_at": now,
            "last_activity": now,
            "pages": pages,
            "total_pages": extracted["total_pages"],
            "chunks": chunks,
            "vector_index": index,
            "parsed_questions": None,
            "validation": None,
            "answer_bank": None,
            "running": False,
            "progress": {"status": "idle", "completed": 0, "total": 0, "failed": 0},
        }
    return {
        "message": "PDF processed and indexed successfully",
        "session_id": session_id,
        "total_pages": extracted["total_pages"],
        "text_pages": extracted.get("digital_pages", 0),
        "ocr_pages": extracted.get("ocr_pages", 0),
        "total_chunks": len(chunks),
        "vector_dimensions": getattr(index, "d", None),
        "characters_extracted": sum(len(str(page.get("text", ""))) for page in pages),
    }


@app.post("/api/sessions/{session_id}/search")
async def search_document(session_id: str, request: SearchRequest):
    session = get_session(session_id)
    matches = await run_in_threadpool(
        search_vector_index, session["vector_index"], session["chunks"], request.question, 5
    )
    return {"session_id": session_id, "matches": matches}


@app.post("/api/sessions/{session_id}/ask")
async def ask_question(session_id: str, request: AskRequest):
    session = get_session(session_id)
    matches = await run_in_threadpool(
        search_vector_index, session["vector_index"], session["chunks"], request.question, 5
    )
    try:
        answer = await run_in_threadpool(
            generate_exam_answer, request.question, matches, None, None, request.allow_general
        )
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Answer service failed: {exc}") from exc
    return {"session_id": session_id, "question": request.question, **answer}


@app.post("/api/sessions/{session_id}/questions/upload")
async def upload_question_paper(session_id: str, file: UploadFile = File(...)):
    session = get_session(session_id)
    if session["running"]:
        raise HTTPException(status_code=409, detail="Answer generation is running.")
    data = await read_pdf(file)
    try:
        parsed = await run_in_threadpool(parse_question_paper_with_vision, data)
        validation = validate_question_paper(parsed)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Question extraction failed: {exc}") from exc
    if not validation["valid"]:
        raise HTTPException(status_code=422, detail={"validation": validation, "parsed": parsed})
    with sessions_lock:
        session["parsed_questions"] = parsed
        session["validation"] = validation
        session["answer_bank"] = None
        session["progress"] = {"status": "idle", "completed": 0, "total": 0, "failed": 0}
    return {"session_id": session_id, "validation": validation, "parsed_questions": parsed}


@app.get("/api/sessions/{session_id}/questions")
def get_questions(session_id: str):
    session = get_session(session_id)
    if session["parsed_questions"] is None:
        raise HTTPException(status_code=404, detail="Upload a question paper first.")
    return {"parsed_questions": session["parsed_questions"], "validation": session["validation"]}


def _generate_in_background(session_id: str):
    with sessions_lock:
        session = sessions_db.get(session_id)
        if not session:
            return
        parsed = session["parsed_questions"]
        index, chunks = session["vector_index"], session["chunks"]
    results = []
    failed = 0
    questions = []
    try:
        questions = flatten_questions(parsed)
        answer_question = create_rag_answer_function(index, chunks)
        for item in questions:
            try:
                if not item["question"].strip():
                    raise ValueError("Empty question text")
                response = answer_question(item["question"], item["marks"])
                if (not isinstance(response, dict)
                    or not str(response.get("answer") or "").strip()
                    or response.get("source") == "insufficient"):
                    raise ValueError("Answer service returned no answer")
                results.append({**item, "status": "completed", **response})
            except ProviderQuotaExceeded as exc:
                # Stop this batch; keep completed answers and mark remaining for retry.
                remaining = questions[len(results):]
                for pending in remaining:
                    results.append({**pending, "status": "failed", "answer": None,
                                    "source": None, "pages": [], "error": str(exc)})
                    failed += 1
                with sessions_lock:
                    session["progress"] = {"status": "running", "completed": len(results),
                                           "total": len(questions), "failed": failed}
                break
            except Exception as exc:
                failed += 1
                results.append({**item, "status": "failed", "answer": None,
                                "source": None, "pages": [], "error": "Answer could not be generated. Retry later."})
            with sessions_lock:
                session["progress"] = {"status": "running", "completed": len(results),
                                       "total": len(questions), "failed": failed}
                session["last_activity"] = utcnow()
        with sessions_lock:
            session["answer_bank"] = {"total_questions": len(questions),
                                      "completed": len(questions) - failed,
                                      "failed": failed, "results": results}
            session["progress"]["status"] = "completed" if failed == 0 else "partial"
    except Exception as exc:
        with sessions_lock:
            session["progress"] = {"status": "failed", "completed": len(results),
                                   "total": len(questions), "failed": failed,
                                   "error": str(exc)}
    finally:
        with sessions_lock:
            session["running"] = False
            session["last_activity"] = utcnow()


@app.post("/api/sessions/{session_id}/answers/generate", status_code=202)
def start_answer_generation(session_id: str):
    session = get_session(session_id)
    with sessions_lock:
        if session["parsed_questions"] is None:
            raise HTTPException(status_code=409, detail="Upload a question paper first.")
        if session["running"]:
            raise HTTPException(status_code=409, detail="Generation already running.")
        session["running"] = True
        session["answer_bank"] = None
        session["progress"] = {"status": "running", "completed": 0,
                               "total": len(flatten_questions(session["parsed_questions"])), "failed": 0}
        threading.Thread(target=_generate_in_background, args=(session_id,), daemon=True).start()
    return {"session_id": session_id, "message": "Generation started", "progress": session["progress"]}


def _retry_failed_in_background(session_id):
    with sessions_lock:
        session = sessions_db.get(session_id)
        if not session:
            return
        bank = session["answer_bank"]
        index, chunks = session["vector_index"], session["chunks"]
    try:
        answer_question = create_rag_answer_function(index, chunks)
        for item in bank["results"]:
            if item.get("status") != "failed":
                continue
            try:
                response = answer_question(item["question"], item.get("marks"))
                if not response.get("answer") or response.get("source") == "insufficient":
                    raise RuntimeError("No usable answer")
                item.update(status="completed", error=None, **response)
            except ProviderQuotaExceeded:
                item["error"] = "Groq quota reached. Retry later."
                break
            except Exception:
                item["error"] = "Answer could not be generated. Retry later."
            with sessions_lock:
                bank["failed"] = sum(x["status"] == "failed" for x in bank["results"])
                bank["completed"] = len(bank["results"]) - bank["failed"]
                session["progress"] = {
                    "status": "running", "completed": bank["completed"],
                    "total": bank["total_questions"], "failed": bank["failed"]
                }
                session["last_activity"] = utcnow()
        with sessions_lock:
            session["progress"]["status"] = "completed" if bank["failed"] == 0 else "partial"
    except Exception as exc:
        with sessions_lock:
            session["progress"] = {"status": "failed", "error": str(exc),
                                   "completed": bank["completed"],
                                   "total": bank["total_questions"], "failed": bank["failed"]}
    finally:
        with sessions_lock:
            session["running"] = False
            session["last_activity"] = utcnow()


@app.post("/api/sessions/{session_id}/answers/retry", status_code=202)
def retry_failed_answers(session_id: str):
    session = get_session(session_id)
    with sessions_lock:
        if session["running"]:
            raise HTTPException(status_code=409, detail="Generation already running")
        bank = session.get("answer_bank")
        if bank is None:
            raise HTTPException(status_code=409, detail="Generate answers first")
        if bank["failed"] == 0:
            return {"message": "No failed answers", "progress": session["progress"]}
        session["running"] = True
        session["progress"]["status"] = "running"
        threading.Thread(target=_retry_failed_in_background,
                         args=(session_id,), daemon=True).start()
    return {"message": "Retry started", "progress": session["progress"]}


@app.get("/api/sessions/{session_id}/answers/progress")
def answer_progress(session_id: str):
    session = get_session(session_id)
    with sessions_lock:
        return dict(session["progress"])


@app.get("/api/sessions/{session_id}/answers")
def get_exam_answers(session_id: str):
    session = get_session(session_id)
    if session["answer_bank"] is None:
        raise HTTPException(status_code=409, detail="Answers not ready yet. Check progress.")
    return session["answer_bank"]


def _export_docx(bank):
    doc = Document()
    doc.add_heading("StudyLens AI - Examination Answer Bank", 0)
    current_section = None
    for item in bank["results"]:
        if item["section"] != current_section:
            current_section = item["section"]
            doc.add_heading(str(current_section), level=1)
        doc.add_heading(f"Q{item['id']} ({item.get('marks') or '?'} marks)", level=2)
        doc.add_paragraph(str(item.get("question") or ""))
        answer = item.get("answer") or "Answer unavailable - retry generation."
        for line in answer.splitlines():
            if line.strip():
                doc.add_paragraph(line.strip())
        if item.get("source") == "general":
            doc.add_paragraph("Source: General AI knowledge (not verified against textbook)")
        elif item.get("source") == "pdf" and item.get("pages"):
            doc.add_paragraph("Textbook pages: " + ", ".join(map(str, item["pages"])))
        if item.get("uncertain"):
            doc.add_paragraph("Note: Original question wording needs review.")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _export_pdf(bank):
    styles = getSampleStyleSheet()
    # Register Unicode fonts for Greek letters and mathematical symbols.
    font_dir = Path(__file__).resolve().parent.parent / "assets" / "fonts"
    system_font = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    win_font = Path("C:/Windows/Fonts/arial.ttf")
    font_file = next((f for f in (font_dir / "DejaVuSans.ttf", system_font, win_font)
                      if f.exists()), None)
    font_name = "Helvetica"
    if font_file:
        try:
            if "StudyLensUnicode" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("StudyLensUnicode", str(font_file)))
            font_name = "StudyLensUnicode"
        except Exception:
            pass
    body = ParagraphStyle("StudyLensBody", parent=styles["BodyText"],
                          fontName=font_name, leading=15, spaceAfter=6)
    heading = ParagraphStyle("StudyLensHeading", parent=styles["Heading2"], fontName=font_name)
    section = ParagraphStyle("StudyLensSection", parent=styles["Heading1"], fontName=font_name)
    title = ParagraphStyle("StudyLensTitle", parent=styles["Title"], fontName=font_name)
    note = ParagraphStyle("StudyLensNote", parent=body, textColor=colors.HexColor("#475569"))
    story = [Paragraph("StudyLens AI - Examination Answer Bank", title), Spacer(1, 12)]
    current_section = None
    for item in bank["results"]:
        if item["section"] != current_section:
            current_section = item["section"]
            story.append(Paragraph(escape(str(current_section)), section))
        story.append(Paragraph(escape(f"Q{item['id']} ({item.get('marks') or '?'} marks)"), heading))
        story.append(Paragraph(escape(str(item.get("question") or "")), body))
        answer = item.get("answer") or "Answer unavailable - retry generation."
        for line in answer.splitlines():
            if line.strip():
                story.append(Paragraph(escape(line.strip()), body))
        if item.get("source") == "general":
            story.append(Paragraph("Source: General AI knowledge (not verified against textbook)", note))
        elif item.get("source") == "pdf" and item.get("pages"):
            story.append(Paragraph("Textbook pages: " + escape(", ".join(map(str, item["pages"]))), note))
        if item.get("uncertain"):
            story.append(Paragraph("Note: Original question wording needs review.", note))
        story.append(Spacer(1, 12))
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=(595.28, 841.89), leftMargin=44,
                      rightMargin=44, topMargin=44, bottomMargin=44).build(story)
    buffer.seek(0)
    return buffer


@app.get("/api/sessions/{session_id}/answers/export/{file_format}")
def export_answer_bank(session_id: str, file_format: str):
    """Return complete downloadable file bytes, not a streaming buffer."""
    session = get_session(session_id)
    with sessions_lock:
        bank = session.get("answer_bank")
        if bank is None:
            raise HTTPException(status_code=409, detail="Generate answers first.")
        # Snapshot results so background retries cannot mutate them mid-export.
        import copy
        bank_snapshot = copy.deepcopy(bank)

    if file_format not in ("docx", "pdf"):
        raise HTTPException(status_code=400, detail="Format must be pdf or docx.")

    try:
        if file_format == "docx":
            buffer = _export_docx(bank_snapshot)
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            expected_magic = b"PK\x03\x04"  # DOCX is a ZIP container
        else:
            buffer = _export_pdf(bank_snapshot)
            media_type = "application/pdf"
            expected_magic = b"%PDF"
        payload = buffer.getvalue()
        if not payload.startswith(expected_magic):
            raise ValueError("Exported file has an invalid format")
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Could not create {file_format.upper()} export: {exc}") from exc

    return Response(
        content=payload,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="studylens_answers.{file_format}"',
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


@app.delete("/api/sessions/{session_id}")
def delete_session(session_id: str):
    with sessions_lock:
        session = sessions_db.get(session_id)
        if session is None:
            raise HTTPException(status_code=404, detail="Session not found.")
        if session["running"]:
            raise HTTPException(status_code=409, detail="Wait for generation to finish before deleting.")
        del sessions_db[session_id]
    return {"message": "Session deleted"}


# Serve the frontend from the same FastAPI process. Keep this LAST.
WEB_DIR = Path(__file__).resolve().parent.parent / "web"
if WEB_DIR.is_dir():
    app.mount("/", StaticFiles(directory=str(WEB_DIR), html=True), name="frontend")
